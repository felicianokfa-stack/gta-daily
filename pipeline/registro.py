"""Cria registro/registro.xlsx: a base de dados do laboratório (Analytics v0).

Uso:
    python -m pipeline.registro            # cria o arquivo se não existir
    python -m pipeline.registro --forcar   # recria do zero (apaga dados!)

Abas: Legenda, Posts, Métricas, Tempo, Radar, Diário.
Métricas e Radar têm fórmulas (taxas por alcance, prioridade), nunca valores fixos.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DEST = ROOT / "registro" / "registro.xlsx"
LINHAS = 300  # linhas pré-formatadas com fórmulas

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
INPUT = Font(name="Arial", size=10, color="0000FF")
HEAD_FILL = PatternFill("solid", fgColor="1E1030")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")


def _sheet(wb, title, headers, widths, calc_cols=()):
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    for r in range(2, LINHAS + 2):
        for i in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=i)
            cell.font = FONT if i in calc_cols else INPUT
            if i in calc_cols:
                cell.fill = CALC_FILL
    return ws


def _lista(ws, col, valores):
    dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{LINHAS + 1}")


def criar(dest: Path):
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------------- Legenda
    lg = wb.create_sheet("Legenda")
    lg.column_dimensions["A"].width = 110
    linhas = [
        ("REGISTRO DO LABORATÓRIO GTA DAILY", True),
        ("Texto azul = campo de preenchimento. Fundo cinza = fórmula (não editar).", False),
        ("Posts: 1 linha por conteúdo (rascunho → aprovado → publicado). O ID é o mesmo da pauta em pautas/.", False),
        ("Métricas: 1 linha por post por janela (24h, 72h, 7d). As taxas por alcance são calculadas.", False),
        ("Tempo: minutos por etapa; 'Quem' separa Agente de Dono (mede onde a automação economiza).", False),
        ("Radar: 1 linha por oportunidade; notas de 1 a 5; a prioridade é a média (critérios em docs/operacao.md).", False),
        ("Diário: decisões, contexto externo (trailer, lançamento) e mudanças de rumo.", False),
        ("", False),
        ("EXEMPLO do formato da aba Métricas (valores fictícios, NÃO são dados do perfil):", True),
        ("2026-09-25-01 | 24h | 2026-09-26 | views 1200 | alcance 950 | curtidas 80 | coment. 12 | compart. 15 | salv. 9 | visitas 30 | seguidores 6", False),
        ("", False),
        ("Regras de evidência: < 5 posts por variante = no máximo SINAL. Ver docs/estado.md.", False),
    ]
    for i, (txt, bold) in enumerate(linhas, start=1):
        c = lg.cell(row=i, column=1, value=txt)
        c.font = Font(name="Arial", size=12 if i == 1 else 10, bold=bold)

    # ---------------------------------------------------------------- Posts
    ph = ["ID", "Data", "Hora", "Série", "Universo", "Tipo info", "Formato", "Duração (s)", "Assunto",
          "Gancho (texto)", "Tipo de gancho", "Estilo de edição", "Origem da ideia", "Fonte", "Origem da mídia",
          "CTA", "Pergunta", "Hipótese testada", "Estado", "Link"]
    pw = [16, 12, 8, 20, 12, 12, 12, 10, 30, 38, 16, 16, 14, 26, 22, 26, 30, 14, 12, 30]
    ps = _sheet(wb, "Posts", ph, pw)
    _lista(ps, "E", ["GTA VI", "GTA V", "GTA Online", "Geral"])
    _lista(ps, "F", ["FATO", "RUMOR", "CURIOSIDADE", "OPINIÃO", "SEMANAL"])
    _lista(ps, "G", ["Reel", "Carrossel", "Reel+Carrossel", "Story", "Post"])
    _lista(ps, "M", ["radar", "atemporal", "impulso"])
    _lista(ps, "S", ["rascunho", "aprovado", "publicado", "descartado"])
    posts = [
        ["2026-09-24-01", None, None, "faltam-x-dias", "GTA VI", "FATO", "Reel+Carrossel", 18, "Contagem + confirmado sobre lançamento",
         "Faltam {dias} dias pro GTA 6", "número", "portal v2", "radar", "Rockstar Newswire, PlayStation Blog", "provisória",
         "Contagem regressiva todo dia", "Você vai jogar no dia do lançamento?", "C6", "rascunho", None],
        ["2026-09-25-01", None, None, None, "GTA VI", "FATO", "Reel+Carrossel", None, "Stephen Root confirmado (Brian Heder)",
         "O primeiro ator do GTA 6 foi confirmado", "fato forte", "portal v2", "radar", "AP via Variety/Kotaku", "provisória",
         "Qual ator você quer no GTA 6?", "Qual ator você quer ver no GTA 6?", "C1", "rascunho", None],
        ["2026-09-25-02", None, None, "gta-online-da-semana", "GTA Online", "SEMANAL", "Reel+Carrossel", None, "Bônus GTA Online 24–30/09",
         "GTA Online: o que vale fazer essa semana", "serviço", "portal v2", "radar", "GTABase, PCQuest, GTA Boss", "provisória",
         "Toda quinta tem", "Você ainda joga GTA Online esperando o 6?", "C2, C4, C6", "rascunho", None],
    ]
    for r, row in enumerate(posts, start=2):
        for c, v in enumerate(row, start=1):
            ps.cell(row=r, column=c, value=v)

    # ---------------------------------------------------------------- Métricas
    mh = ["Post ID", "Janela", "Data coleta", "Views", "Alcance", "Curtidas", "Comentários", "Compartilhamentos",
          "Salvamentos", "Visitas ao perfil", "Seguidores ganhos", "Tempo médio assistido (s)", "Taxa de pulo",
          "Compart./alcance", "Salv./alcance", "Seguidores/alcance", "Interações/alcance", "Visitas/alcance"]
    mw = [16, 9, 12, 10, 10, 10, 12, 16, 12, 14, 14, 16, 12, 14, 13, 16, 16, 13]
    ms = _sheet(wb, "Métricas", mh, mw, calc_cols=range(14, 19))
    _lista(ms, "B", ["24h", "72h", "7d"])
    for r in range(2, LINHAS + 2):
        guard = f'IF(OR($E{r}="",$E{r}=0),"",'
        ms.cell(row=r, column=14, value=f"={guard}H{r}/E{r})")
        ms.cell(row=r, column=15, value=f"={guard}I{r}/E{r})")
        ms.cell(row=r, column=16, value=f"={guard}K{r}/E{r})")
        ms.cell(row=r, column=17, value=f"={guard}(F{r}+G{r}+H{r}+I{r})/E{r})")
        ms.cell(row=r, column=18, value=f"={guard}J{r}/E{r})")
        for c in range(14, 19):
            ms.cell(row=r, column=c).number_format = "0.0%"
        ms.cell(row=r, column=13).number_format = "0.0%"
    ms["N1"].comment = Comment("Taxas = métrica / alcance. Em branco enquanto não houver alcance.", "Agente")

    # ---------------------------------------------------------------- Tempo
    th = ["Post ID", "Quem", "Pesquisa", "Ideia", "Roteiro", "Busca de material", "Edição", "Legenda",
          "Publicação", "Revisão/aprovação", "Registro", "Total (min)", "Observação"]
    tw = [16, 10, 10, 8, 9, 14, 8, 9, 11, 14, 9, 11, 40]
    ts = _sheet(wb, "Tempo", th, tw, calc_cols=[12])
    _lista(ts, "B", ["Agente", "Dono"])
    for r in range(2, LINHAS + 2):
        ts.cell(row=r, column=12, value=f'=IF(COUNT(C{r}:K{r})=0,"",SUM(C{r}:K{r}))')
    tempos = [
        ["2026-09-25-01", "Agente", 8, 2, 3, 0, 1, 1, None, None, 1, None, "Estimativa do Agente; edição = render automático (~1 min)"],
        ["2026-09-25-02", "Agente", 6, 1, 3, 0, 1, 1, None, None, 1, None, "Estimativa do Agente; radar compartilhado com 25-01"],
    ]
    for r, row in enumerate(tempos, start=2):
        for c, v in enumerate(row, start=1):
            if c != 12:
                ts.cell(row=r, column=c, value=v)

    # ---------------------------------------------------------------- Radar
    rh = ["Data", "#", "Oportunidade", "Universo", "Status info", "Relevância", "Momento", "Diferenciação",
          "Viabilidade", "Valor de aprendizado", "Prioridade", "Decisão", "Pauta"]
    rw = [12, 5, 46, 12, 12, 11, 10, 13, 11, 13, 11, 34, 16]
    rs = _sheet(wb, "Radar", rh, rw, calc_cols=[11])
    for r in range(2, LINHAS + 2):
        rs.cell(row=r, column=11, value=f'=IF(COUNT(F{r}:J{r})=0,"",AVERAGE(F{r}:J{r}))')
        rs.cell(row=r, column=11).number_format = "0.0"
    radar = [
        (1, "Stephen Root é o primeiro ator confirmado do GTA 6", "GTA VI", "FATO", 5, 5, 4, 4, 4, "Pauta hoje", "2026-09-25-01"),
        (2, "GTA Online: o que vale fazer em 24–30/09", "GTA Online", "REPORTADO", 4, 5, 3, 4, 5, "Pauta hoje (série semanal)", "2026-09-25-02"),
        (3, "Faltam X dias: contagem diária", "GTA VI", "FATO", 4, 3, 3, 5, 4, "Série diária", "2026-09-24-01"),
        (4, "Camisa do Tommy Vercetti no bônus de pré-venda", "GTA VI", "REPORTADO", 4, 3, 5, 2, 5, "Verificar no Newswire", None),
        (5, "10 detalhes do Extended Look (Netflix 27/08)", "GTA VI", "FATO", 5, 3, 4, 1, 5, "Depende de mídia oficial", None),
        (6, "Relacionamento Jason × Lucia", "GTA VI", "REPORTADO", 4, 3, 3, 3, 3, "Guardar; enquete em Stories", None),
        (7, "Que horas o GTA 6 libera no Brasil?", "GTA VI", "REPORTADO", 5, 2, 4, 2, 4, "Pesquisar perto do lançamento", None),
        (8, "DualSense do GTA 6 (pré-venda 10/09)", "GTA VI", "FATO", 3, 1, 2, 4, 2, "Descartado (saturado)", None),
    ]
    for r, (n, op, uni, st, a, b, c, d, e, dec, pauta) in enumerate(radar, start=2):
        for col, v in zip([1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 12, 13],
                          ["2026-09-25", n, op, uni, st, a, b, c, d, e, dec, pauta]):
            rs.cell(row=r, column=col, value=v)

    # ---------------------------------------------------------------- Diário
    dh = ["Data", "Tipo", "Descrição"]
    ds = _sheet(wb, "Diário", dh, [12, 14, 110])
    _lista(ds, "B", ["decisão", "contexto", "evento", "correção", "feedback"])
    diario = [
        ("2026-09-24", "decisão", "Material: texto + mídia oficial com crédito; sem vazamento, sem corte sem autorização."),
        ("2026-09-24", "feedback", "Criativo v1 reprovado (sem imagem, genérico). Adotado o estilo 'portal de notícias gamer' (v2)."),
        ("2026-09-25", "decisão", "Novo documento-base: IA × Agente, operação proativa, escopo em todo o universo GTA."),
        ("2026-09-25", "correção", "DualSense do GTA 6: revelado em 03/09, pré-venda em 10/09 (antes estava 'mostrou em 10/09')."),
        ("2026-09-25", "contexto", "Faltam 55 dias para o GTA VI (19/11/2026). Rede bloqueia rockstargames.com."),
    ]
    for r, row in enumerate(diario, start=2):
        for c, v in enumerate(row, start=1):
            ds.cell(row=r, column=c, value=v)

    # sem valores em cache (openpyxl): Excel/Sheets recalculam tudo ao abrir
    wb.calculation.fullCalcOnLoad = True
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--forcar", action="store_true")
    args = ap.parse_args()
    if DEST.exists() and not args.forcar:
        print(f"{DEST} já existe (use --forcar para recriar e apagar os dados)")
        return
    print(criar(DEST))


if __name__ == "__main__":
    main()
